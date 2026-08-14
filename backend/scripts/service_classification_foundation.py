import json
import os
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import text

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from app.database import engine, SessionLocal
from app.models import Transaction, ServiceClassification


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_XLSX_CANDIDATES = [
    PROJECT_ROOT / "DM_DV.xlsx",
    PROJECT_ROOT / "backend" / "data" / "DM_DV.xlsx",
    PROJECT_ROOT / "data" / "DM_DV.xlsx",
]
REPORT_DIR = PROJECT_ROOT / "backend" / "data" / "reports"
REPORT_DIR.mkdir(parents=True, exist_ok=True)


def find_taxonomy_file():
    for candidate in DEFAULT_XLSX_CANDIDATES:
        if candidate.exists():
            return candidate
    raise FileNotFoundError("Không tìm thấy DM_DV.xlsx trong các vị trí dự kiến.")


def normalize(value):
    if value is None:
        return ""
    return str(value).strip()


def resolve_columns(headers):
    idx_type = idx_code = idx_name = idx_note = None
    for i, header in enumerate(headers):
        h = normalize(header).lower()
        if idx_type is None and ("loại dịch vụ" in h or "loai dich vu" in h or "loaidichvu" in h):
            idx_type = i
        elif idx_code is None and ("dichvuchinh" in h or "dịch vụ chính" in h or "dich vu chinh" in h):
            idx_code = i
        elif idx_name is None and ("tên dịch vụ" in h or "ten dich vu" in h or "tên dv" in h):
            idx_name = i
        elif idx_note is None and ("ghi chú" in h or "ghi chu" in h):
            idx_note = i
    if idx_type is None or idx_code is None:
        raise ValueError(f"Không xác định được cột taxonomy từ headers: {headers}")
    return idx_type, idx_code, idx_name, idx_note


def load_taxonomy(xlsx_path: Path):
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx_type, idx_code, idx_name, idx_note = resolve_columns(headers)

    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in r):
            continue
        loai = normalize(r[idx_type])
        ma = normalize(r[idx_code]).upper()
        ten = normalize(r[idx_name]) if idx_name is not None else ""
        note = normalize(r[idx_note]) if idx_note is not None else ""
        if not ma:
            continue
        rows.append({
            "ma_dv": ma,
            "loai_dich_vu": loai or "Khác",
            "dich_vu_chinh": ten,
            "ghi_chu": note,
        })
    return rows


def ensure_schema(conn):
    conn.execute(text("""
        CREATE TABLE IF NOT EXISTS service_classifications (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            ma_dv VARCHAR(50) NOT NULL UNIQUE,
            loai_dich_vu VARCHAR(100) NOT NULL,
            is_active BOOLEAN DEFAULT 1,
            created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
            updated_at DATETIME DEFAULT CURRENT_TIMESTAMP
        )
    """))
    conn.execute(text("CREATE INDEX IF NOT EXISTS idx_service_classifications_ma_dv ON service_classifications(ma_dv)"))
    conn.execute(text("CREATE INDEX IF NOT EXISTS idx_service_classifications_loai_dich_vu ON service_classifications(loai_dich_vu)"))
    conn.execute(text("CREATE INDEX IF NOT EXISTS idx_service_classifications_is_active ON service_classifications(is_active)"))

    # Add transactions.loai_dich_vu if missing
    cols = conn.execute(text("PRAGMA table_info(transactions)")).fetchall()
    col_names = {row[1] for row in cols}
    if "loai_dich_vu" not in col_names:
        conn.execute(text("ALTER TABLE transactions ADD COLUMN loai_dich_vu VARCHAR(100)"))
    conn.execute(text("CREATE INDEX IF NOT EXISTS idx_trans_loai_dich_vu ON transactions(loai_dich_vu)"))


def seed_taxonomy(conn, taxonomy_rows):
    inserted = 0
    updated = 0
    for row in taxonomy_rows:
        existing = conn.execute(
            text("SELECT id, loai_dich_vu, is_active FROM service_classifications WHERE ma_dv = :ma_dv"),
            {"ma_dv": row["ma_dv"]},
        ).fetchone()
        if existing:
            conn.execute(
                text("""
                    UPDATE service_classifications
                    SET loai_dich_vu = :loai_dich_vu,
                        is_active = 1,
                        updated_at = CURRENT_TIMESTAMP
                    WHERE ma_dv = :ma_dv
                """),
                {"ma_dv": row["ma_dv"], "loai_dich_vu": row["loai_dich_vu"]},
            )
            updated += 1
        else:
            conn.execute(
                text("""
                    INSERT INTO service_classifications (ma_dv, loai_dich_vu, is_active)
                    VALUES (:ma_dv, :loai_dich_vu, 1)
                """),
                {"ma_dv": row["ma_dv"], "loai_dich_vu": row["loai_dich_vu"]},
            )
            inserted += 1
    return inserted, updated


def backfill_transactions(conn):
    mapping = {
        row[0]: row[1]
        for row in conn.execute(text("""
            SELECT ma_dv, loai_dich_vu
            FROM service_classifications
            WHERE is_active = 1
        """)).fetchall()
    }
    tx_rows = conn.execute(text("SELECT id, ma_dv, dich_vu_chinh FROM transactions")).fetchall()
    conn.commit()
    total = len(tx_rows)
    mapping_success = 0
    fallback = 0
    missing_codes = Counter()

    updates = []
    for tx_id, ma_dv, dich_vu_chinh in tx_rows:
        key = normalize(dich_vu_chinh).upper()
        loai = mapping.get(key)
        if loai:
            mapping_success += 1
        else:
            loai = "Khác"
            fallback += 1
            missing_codes[key or "NULL"] += 1
        updates.append({"id": tx_id, "loai_dich_vu": loai})

    if updates:
        with conn.begin():
            conn.execute(text("UPDATE transactions SET loai_dich_vu = :loai_dich_vu WHERE id = :id"), updates)

    top10 = conn.execute(text("""
        SELECT ma_dv, COUNT(*) AS total
        FROM transactions
        GROUP BY ma_dv
        ORDER BY total DESC, ma_dv ASC
        LIMIT 10
    """)).fetchall()

    return {
        "total_transactions": total,
        "mapping_success": mapping_success,
        "fallback_khac": fallback,
        "missing_codes": missing_codes,
        "top10_ma_dv": [{"ma_dv": r[0], "count": r[1]} for r in top10],
    }


def validation_report(conn, seed_rows, backfill_result):
    taxonomy_counts = Counter(r["loai_dich_vu"] for r in seed_rows)
    distinct_taxonomy = sorted(taxonomy_counts.keys())
    service_rows = conn.execute(text("""
        SELECT ma_dv, loai_dich_vu, COUNT(*) AS cnt
        FROM service_classifications
        GROUP BY ma_dv, loai_dich_vu
        ORDER BY ma_dv
    """)).fetchall()

    report = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "taxonomy_distinct": distinct_taxonomy,
        "taxonomy_counts": dict(sorted(taxonomy_counts.items(), key=lambda x: x[0])),
        "service_classifications": [
            {"ma_dv": r[0], "loai_dich_vu": r[1], "count": r[2]} for r in service_rows
        ],
        "backfill": backfill_result,
    }
    report_path = REPORT_DIR / f"service_classification_validation_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report, report_path


def main():
    xlsx_path = find_taxonomy_file()
    print(f"[FOUNDATION] DM_DV.xlsx: {xlsx_path}")
    taxonomy_rows = load_taxonomy(xlsx_path)
    print(f"[FOUNDATION] Seed rows loaded: {len(taxonomy_rows)}")

    conn = engine.connect()
    try:
        ensure_schema(conn)
        print("[FOUNDATION] Schema ensured.")

        inserted, updated = seed_taxonomy(conn, taxonomy_rows)
        print(f"[FOUNDATION] Seed complete. inserted={inserted}, updated={updated}")
        conn.commit()

        backfill_result = backfill_transactions(conn)
        print("[FOUNDATION] Backfill complete.")

        report, report_path = validation_report(conn, taxonomy_rows, backfill_result)
        print(f"[FOUNDATION] Validation report saved: {report_path}")

        print(json.dumps({
            "root_cause": "Thiếu bảng taxonomy riêng và transaction chưa có cột classification layer.",
            "migration_applied": [
                "CREATE TABLE IF NOT EXISTS service_classifications",
                "ALTER TABLE transactions ADD COLUMN loai_dich_vu",
            ],
            "backfill_result": backfill_result,
        }, ensure_ascii=False, indent=2))
    finally:
        conn.close()


if __name__ == "__main__":
    main()
